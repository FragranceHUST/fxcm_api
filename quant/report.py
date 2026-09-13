"""Excel 报告输出：openpyxl 追加式工作簿（数据表 + meta 流水），多批结果共存于同一文件。"""
from __future__ import annotations

import json
import math
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

META_SHEET = "meta"
MAX_TITLE = 31


def _cell(v):
    """None/非有限浮点 → 空单元格；dict → JSON 字符串；其余原样。"""
    if v is None:
        return None
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def _unique_title(wb: Workbook, title: str) -> str:
    """31 字符安全去重：冲突时追加 -2、-3……后缀并同步截短基底。"""
    base = title[:MAX_TITLE]
    candidate = base
    k = 2
    while candidate in wb.sheetnames:
        suffix = f"-{k}"
        candidate = base[:MAX_TITLE - len(suffix)] + suffix
        k += 1
    return candidate


def append_sheet(xlsx_path: str | Path, title: str, headers: list[str],
                 rows: list[list], meta_lines: list[str] | None = None) -> None:
    """向工作簿追加一张数据表（永不覆盖已有 sheet）；meta_lines 追加到 meta 流水表。"""
    path = Path(xlsx_path)
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        first = wb.active
        assert first is not None
        wb.remove(first)
        wb.create_sheet(META_SHEET)
    if META_SHEET not in wb.sheetnames:
        wb.create_sheet(META_SHEET)

    ws = wb.create_sheet(_unique_title(wb, title))
    ws.append([_cell(h) for h in headers])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in rows:
        ws.append([_cell(v) for v in row])
    ws.freeze_panes = "A2"

    if meta_lines:
        ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        wb[META_SHEET].append([ts, *[_cell(line) for line in meta_lines]])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
