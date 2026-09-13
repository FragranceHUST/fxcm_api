"""report.append_sheet 测试：建簿、追加去重、meta 流水、类型归一（全离线）。"""
from __future__ import annotations

import tempfile
import unittest

from openpyxl import load_workbook

from quant.report import append_sheet


class TestAppendSheet(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = f"{self.tmp.name}/results.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def test_create_then_append_with_dedup(self):
        append_sheet(self.path, "long_sweep", ["a", "b"],
                     [[1, 1.5], [2, None]], meta_lines=["run1"])
        append_sheet(self.path, "long_sweep", ["a", "b"],
                     [[3, 2.5]], meta_lines=["run2"])
        wb = load_workbook(self.path)
        self.assertEqual(wb.sheetnames, ["meta", "long_sweep", "long_sweep-2"])
        ws = wb["long_sweep"]
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.cell(1, 1).value, "a")
        self.assertEqual(ws.cell(2, 1).value, 1)
        self.assertEqual(ws.cell(2, 2).value, 1.5)
        self.assertIsNone(ws.cell(3, 2).value)
        ws2 = wb["long_sweep-2"]
        self.assertEqual(ws2.max_row, 2)
        self.assertEqual(ws2.cell(2, 1).value, 3)
        meta = wb["meta"]
        self.assertEqual(meta.max_row, 2)
        self.assertEqual(meta.cell(1, 2).value, "run1")
        self.assertEqual(meta.cell(2, 2).value, "run2")
        self.assertIsInstance(meta.cell(1, 1).value, str)

    def test_header_bold_and_freeze(self):
        append_sheet(self.path, "t", ["h1", "h2"], [[1, 2]])
        ws = load_workbook(self.path)["t"]
        self.assertTrue(ws.cell(1, 1).font.bold)
        self.assertEqual(ws.freeze_panes, "A2")

    def test_none_empty_float_number_dict_json(self):
        append_sheet(self.path, "t", ["k", "f", "d"],
                     [[1, 2.5, {"2020": 1.5}], [None, None, None]])
        ws = load_workbook(self.path)["t"]
        self.assertEqual(ws.cell(2, 2).value, 2.5)
        self.assertIsInstance(ws.cell(2, 2).value, float)
        self.assertEqual(ws.cell(2, 3).value, '{"2020": 1.5}')
        self.assertIsNone(ws.cell(3, 1).value)
        self.assertIsNone(ws.cell(3, 3).value)

    def test_title_truncated_to_31_and_dedup(self):
        long_title = "long_" + "x" * 40
        append_sheet(self.path, long_title, ["a"], [[1]])
        append_sheet(self.path, long_title, ["a"], [[2]])
        wb = load_workbook(self.path)
        names = [n for n in wb.sheetnames if n != "meta"]
        self.assertEqual(len(names), 2)
        for n in names:
            self.assertLessEqual(len(n), 31)
        self.assertEqual(names[1], names[0][:29] + "-2")

    def test_sorted_rows_pass_through(self):
        rows = [[3], [1], [2]]
        append_sheet(self.path, "t", ["v"], rows)
        ws = load_workbook(self.path)["t"]
        self.assertEqual([ws.cell(r, 1).value for r in (2, 3, 4)], [3, 1, 2])

    def test_no_meta_lines_keeps_meta_empty(self):
        append_sheet(self.path, "t", ["a"], [[1]])
        wb = load_workbook(self.path)
        self.assertIn("meta", wb.sheetnames)
        self.assertEqual(wb["meta"].max_row, 1)


if __name__ == "__main__":
    unittest.main()
